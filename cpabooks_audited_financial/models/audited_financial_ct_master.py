# -*- coding: utf-8 -*-
import base64
import io
import re
from copy import copy
from datetime import date

from odoo import fields, models, _
from odoo.exceptions import UserError


class AuditedFinancialVersionCtMaster(models.Model):
    _inherit = "audited.financial.version"

    def _ct_master_company(self):
        self.ensure_one()
        companies = self.company_ids or self.env.company
        return companies[:1]

    def _ct_master_company_label(self):
        return (self._ct_master_company().name or "Company").strip()

    def _ct_master_year(self):
        try:
            return int(self.year_current or 0) or date.today().year
        except (TypeError, ValueError):
            return date.today().year

    def _ct_master_safe_name(self, name):
        cleaned = re.sub(r'[\\/*?:\[\]]', " ", name or "file")
        cleaned = re.sub(r"\s+", " ", cleaned).strip()
        return cleaned[:120] or "file"

    def _ct_master_attachment(self, filename, data, mimetype):
        att = self.env["ir.attachment"].create({
            "name": filename,
            "type": "binary",
            "datas": base64.b64encode(data),
            "res_model": self._name,
            "res_id": self.id,
            "mimetype": mimetype,
        })
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=true" % att.id,
            "target": "self",
        }

    def _ct_master_collect_xlsx(self, files, skipped, filename, builder):
        try:
            data = builder()
            if data:
                files.append((filename, data))
            else:
                skipped.append("%s — no data" % filename)
        except Exception as err:
            skipped.append("%s — %s" % (filename, err))

    def _ct_master_unique_sheet(self, used, name):
        base = re.sub(r'[\\/*?:\[\]]', "", (name or "Sheet"))[:31] or "Sheet"
        candidate = base
        i = 2
        while candidate.lower() in used:
            suffix = "_%s" % i
            candidate = (base[: 31 - len(suffix)] + suffix)[:31]
            i += 1
        used.add(candidate.lower())
        return candidate

    def _ct_master_copy_sheet(self, src_ws, dest_wb, title):
        dst = dest_wb.create_sheet(title)
        for row in src_ws.iter_rows():
            for cell in row:
                new_cell = dst.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.fill = copy(cell.fill)
                    new_cell.border = copy(cell.border)
                    new_cell.alignment = copy(cell.alignment)
                    new_cell.number_format = cell.number_format
        for merged in list(src_ws.merged_cells.ranges):
            try:
                dst.merge_cells(str(merged))
            except Exception:
                pass
        for col_letter, dim in src_ws.column_dimensions.items():
            if dim.width:
                dst.column_dimensions[col_letter].width = dim.width
        for row_idx, dim in src_ws.row_dimensions.items():
            if dim.height:
                dst.row_dimensions[row_idx].height = dim.height
        try:
            dst.freeze_panes = src_ws.freeze_panes
            dst.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
            dst.page_setup.orientation = src_ws.page_setup.orientation
        except Exception:
            pass
        return dst

    def _ct_master_merge_workbooks(self, files, skipped):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            raise UserError(_(
                "CT Filing Master Files needs the openpyxl Python package on the server."
            ))
        dest = Workbook()
        default = dest.active
        dest.remove(default)
        used = set()
        for prefix, data in files:
            src = load_workbook(io.BytesIO(data), data_only=False)
            sheets = src.worksheets
            for ws in sheets:
                if len(sheets) == 1:
                    title = self._ct_master_unique_sheet(used, prefix)
                else:
                    title = self._ct_master_unique_sheet(
                        used, "%s %s" % (prefix, ws.title or "Sheet")
                    )
                self._ct_master_copy_sheet(ws, dest, title)
            src.close()
        if skipped:
            notes = dest.create_sheet(self._ct_master_unique_sheet(used, "Skipped"))
            notes["A1"] = "Not included"
            row = 2
            for line in skipped:
                notes.cell(row=row, column=1, value=line)
                row += 1
        if not dest.worksheets:
            raise UserError(_("No Excel sheets could be built for CT Filing Master Files."))
        buf = io.BytesIO()
        dest.save(buf)
        return buf.getvalue()

    def _ct_master_l1_xlsx_bytes(self, max_level):
        result = self.with_context(afg_return_xlsx_bytes=True).action_export_audit_xlsx(
            max_level=max_level
        )
        if isinstance(result, (bytes, bytearray)):
            return bytes(result)
        if isinstance(result, dict):
            raise UserError(
                result.get("name")
                or _("Complete L1 print setup / gate before CT Filing Master Files.")
            )
        return b""

    def _ct_master_ctc_xlsx(self, company, year):
        if "ctc.report.formula" not in self.env:
            return b""
        CTC = self.env["ctc.report.formula"].sudo()
        rec = CTC.search([("scope_company_ids", "in", company.ids)], limit=1)
        if not rec:
            rec = CTC.search([], limit=1)
        if not rec or not hasattr(rec, "_ctc_employees_salaries_eosb_bytes"):
            return b""
        return rec.with_context(allowed_company_ids=company.ids)._ctc_employees_salaries_eosb_bytes(
            year=year
        )

    def _ct_master_rent_xlsx(self, company, year):
        if "rent.expense.prepaid.wizard" not in self.env:
            return b""
        wiz = self.env["rent.expense.prepaid.wizard"].sudo().with_company(company).create({
            "year": str(year),
        })
        if hasattr(wiz, "_rent_prepaid_xlsx_bytes"):
            return wiz._rent_prepaid_xlsx_bytes()
        return b""

    def _ct_master_far_xlsx(self, company, year, schedule_type):
        if "mej.far.schedule" not in self.env:
            return b""
        Far = self.env["mej.far.schedule"].sudo().with_company(company)
        rec = Far.search([
            ("schedule_type", "=", schedule_type),
            ("company_id", "=", company.id),
            ("year", "=", year),
        ], limit=1)
        if not rec:
            rec = Far.search([
                ("schedule_type", "=", schedule_type),
                ("company_id", "=", company.id),
            ], limit=1)
        if not rec:
            rec = Far.with_context(default_company_id=company.id)._get_or_create(schedule_type)
            if rec.year != year:
                rec.year = year
        if hasattr(rec, "_far_report_xlsx_bytes"):
            return rec._far_report_xlsx_bytes()
        return b""

    def _ct_master_afg_dates(self):
        self.ensure_one()
        df_c = fields.Date.to_string(self.date_from_current) if self.date_from_current else "%s-01-01" % self._ct_master_year()
        dt_c = fields.Date.to_string(self.date_to_current) if self.date_to_current else "%s-12-31" % self._ct_master_year()
        yp = self.year_prior or (self._ct_master_year() - 1)
        df_p = fields.Date.to_string(self.date_from_prior) if self.date_from_prior else "%s-01-01" % yp
        dt_p = fields.Date.to_string(self.date_to_prior) if self.date_to_prior else "%s-12-31" % yp
        return df_c, dt_c, df_p, dt_p

    def _ct_master_report_options(self, view_mode="default", date_mode="range"):
        df_c, dt_c, df_p, dt_p = self._ct_master_afg_dates()
        prior = {
            "string": str(self.year_prior or dt_p[:4]),
            "period_type": "fiscalyear",
            "mode": date_mode,
            "date_from": df_p,
            "date_to": dt_p,
            "filter": "custom",
        }
        if date_mode == "single":
            prior["date"] = dt_p
        date_opts = {
            "date_from": df_c,
            "date_to": dt_c,
            "filter": "custom",
            "mode": date_mode,
            "string": str(self.year_current or dt_c[:4]),
        }
        if date_mode == "single":
            date_opts["date"] = dt_c
        return {
            "date": date_opts,
            "comparison": {
                "filter": "custom",
                "number_period": 1,
                "date_from": df_p,
                "date_to": dt_p,
                "periods": [prior],
            },
            "unfold_all": True,
            "hierarchy": False,
            "all_entries": False,
            "cpa_hide_zero": False,
            "cpa_hide_codes": False,
            "cpa_pl_view_mode": view_mode,
            "cpa_pl_newest_first": False,
        }

    def _ct_master_apply_afg_periods(self, options, date_mode="range"):
        df_c, dt_c, df_p, dt_p = self._ct_master_afg_dates()
        prior = {
            "string": str(self.year_prior or dt_p[:4]),
            "period_type": "fiscalyear",
            "mode": date_mode,
            "date_from": df_p,
            "date_to": dt_p,
            "filter": "custom",
        }
        if date_mode == "single":
            prior["date"] = dt_p
        date_opts = dict(options.get("date") or {})
        date_opts.update({
            "date_from": df_c,
            "date_to": dt_c,
            "filter": "custom",
            "mode": date_mode,
            "string": str(self.year_current or dt_c[:4]),
        })
        if date_mode == "single":
            date_opts["date"] = dt_c
        options["date"] = date_opts
        options["comparison"] = {
            "filter": "custom",
            "number_period": 1,
            "date_from": df_p,
            "date_to": dt_p,
            "periods": [prior],
        }
        options["unfold_all"] = True
        return options

    def _ct_master_fin_report_xlsx(self, xmlids, view_mode="default", date_mode="range"):
        report = False
        for xmlid in xmlids:
            report = self.env.ref(xmlid, raise_if_not_found=False)
            if report:
                break
        if not report:
            return b""
        prev = self._ct_master_report_options(view_mode=view_mode, date_mode=date_mode)
        options = report._get_options(prev)
        options = self._ct_master_apply_afg_periods(options, date_mode=date_mode)
        options["cpa_pl_view_mode"] = view_mode
        options["unfold_all"] = True
        options["hierarchy"] = False
        data = report.get_xlsx(options)
        return data or b""

    def _ct_master_tb_xlsx(self):
        if "account.coa.report" not in self.env:
            return b""
        report = self.env["account.coa.report"]
        prev = self._ct_master_report_options(view_mode="type", date_mode="range")
        options = report._get_options(prev)
        options = self._ct_master_apply_afg_periods(options, date_mode="range")
        options["unfold_all"] = True
        options["cpa_pl_view_mode"] = "type"
        options["hierarchy"] = False
        options["cpa_tb_by_ledger_type"] = True
        if hasattr(report, "get_xlsx"):
            return report.get_xlsx(options) or b""
        return b""

    def action_export_ct_filing_master(self, max_level=None):
        self.ensure_one()
        company = self._ct_master_company()
        year = self._ct_master_year()
        label = self._ct_master_company_label()
        xlsx_name = self._ct_master_safe_name(
            "CT Filing Master Files - %s" % label
        ) + ".xlsx"
        files = []
        skipped = []
        l1 = self.with_context(afg_return_xlsx_bytes=True).action_export_audit_xlsx(
            max_level=max_level
        )
        if isinstance(l1, dict):
            return l1
        if l1:
            files.append(("01 L1 Official", bytes(l1)))
        else:
            skipped.append("01 L1 Official — empty")

        self._ct_master_collect_xlsx(
            files, skipped,
            "02 Salaries EOSB",
            lambda: self._ct_master_ctc_xlsx(company, year),
        )
        self._ct_master_collect_xlsx(
            files, skipped,
            "03 Rent Prepaid",
            lambda: self._ct_master_rent_xlsx(company, year),
        )
        self._ct_master_collect_xlsx(
            files, skipped,
            "04 FAR Auto",
            lambda: self._ct_master_far_xlsx(company, year, "books"),
        )
        self._ct_master_collect_xlsx(
            files, skipped,
            "05 FAR Manual",
            lambda: self._ct_master_far_xlsx(company, year, "manual"),
        )
        self._ct_master_collect_xlsx(
            files, skipped,
            "06 Profit and Loss",
            lambda: self._ct_master_fin_report_xlsx(
                (
                    "account_reports.account_financial_report_profitandloss0",
                    "account_reports.financial_report_profit_and_loss",
                ),
                view_mode="default",
                date_mode="range",
            ),
        )
        self._ct_master_collect_xlsx(
            files, skipped,
            "07 Balance Sheet",
            lambda: self._ct_master_fin_report_xlsx(
                (
                    "account_reports.account_financial_report_balancesheet0",
                    "account_reports.financial_report_balance_sheet",
                ),
                view_mode="default",
                date_mode="single",
            ),
        )
        self._ct_master_collect_xlsx(
            files, skipped,
            "08 Trial Balance",
            lambda: self._ct_master_tb_xlsx(),
        )

        if not files:
            raise UserError(_("No Excel files could be built for CT Filing Master Files."))

        data = self._ct_master_merge_workbooks(files, skipped)
        return self._ct_master_attachment(
            xlsx_name,
            data,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
