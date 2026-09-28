# -*- coding: utf-8 -*-

import base64
import io
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:  # pragma: no cover - tenant images may omit Excel extras
    openpyxl = None
    Font = Alignment = PatternFill = Border = Side = None

from odoo import models, fields, api, _
from odoo.exceptions import UserError


class CafmAmcInvoicePrintWizard(models.TransientModel):
    _name = 'cpabooks.amc.invoice.print.wizard'
    _description = 'AMC Invoice Print / Export Wizard'

    partner_id = fields.Many2one('res.partner', string='Client Name')
    as_of_date = fields.Date(string='As of Date', default=fields.Date.context_today, required=True)
    status = fields.Selection(
        [
            ('waiting_kpi', 'Waiting for KPI'),
            ('waiting_csr', 'Waiting for CSR'),
            ('not_submitted', 'Not Submitted'),
            ('submitted', 'Submitted'),
        ],
        string='Status',
        help='Leave empty to include all statuses.',
    )
    follow_up_user_id = fields.Many2one(
        'res.users',
        string='Follow up by',
        help='Leave empty to include all follow-up users.',
    )

    col_invoice_date = fields.Boolean(string='Inv. Date', default=True)
    col_ref_no = fields.Boolean(string='ref. no (tally)', default=True)
    col_project_name = fields.Boolean(string='Project Name', default=True)
    col_party_name = fields.Boolean(string='Client Name', default=True)
    col_l2_level = fields.Boolean(string='L2 Level (Client Group)', default=True)
    col_l3_level = fields.Boolean(string='L3 Level (Project)', default=True)
    col_invoice_amount = fields.Boolean(string='Inv. Amount', default=True)
    col_ubs_no = fields.Boolean(string='UBS No.', default=True)
    col_contract_lpo_no = fields.Boolean(string='Contract/LPO No.', default=True)
    col_due_day = fields.Boolean(string='Due Days', default=True)
    col_invoice_type = fields.Boolean(string='Inv. Type', default=True)
    col_follow_up = fields.Boolean(string='Follow up by', default=True)
    col_remarks = fields.Boolean(string='Remarks', default=True)
    col_status = fields.Boolean(string='Status', default=True)

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        last_wizard = self.search([('create_uid', '=', self.env.uid)], order='id desc', limit=1)
        if last_wizard:
            for col in [f for f in self._fields if f.startswith('col_')]:
                if col in fields_list:
                    res[col] = getattr(last_wizard, col)
        return res

    def _get_filtered_records(self):
        domain = [('invoice_date', '<=', self.as_of_date)]
        if self.partner_id:
            domain.append(('partner_id', '=', self.partner_id.id))
        if self.status:
            domain.append(('status', '=', self.status))
        if self.follow_up_user_id:
            name = (self.follow_up_user_id.name or '').strip()
            domain += [
                '|',
                ('follow_up_user_id', '=', self.follow_up_user_id.id),
                ('follow_up', 'ilike', name),
            ]
        return self.env['cpabooks.cafm.amc.invoice.reg'].search(
            domain, order='partner_id, invoice_date asc, id asc',
        )

    def _get_grouped_data(self, records):
        # Client → Year → Month (matches on-screen list clubbing)
        by_client = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
        for r in records:
            if not r.invoice_date:
                continue
            client_key = r.partner_id.display_name if r.partner_id else (r.project_name or 'No Client')
            year = r.invoice_date.year
            month = r.invoice_date.strftime('%B %Y').upper()
            by_client[client_key][year][month].append(r)

        month_order = {
            'JANUARY': 1, 'FEBRUARY': 2, 'MARCH': 3, 'APRIL': 4,
            'MAY': 5, 'JUNE': 6, 'JULY': 7, 'AUGUST': 8,
            'SEPTEMBER': 9, 'OCTOBER': 10, 'NOVEMBER': 11, 'DECEMBER': 12,
        }
        res = []
        for client in sorted(by_client.keys(), key=lambda n: (n or '').lower()):
            years = []
            for year in sorted(by_client[client].keys()):
                months = by_client[client][year]
                month_list = []
                for m in sorted(months.keys(), key=lambda name: month_order.get(name.split()[0], 0)):
                    month_list.append({'name': m, 'records': months[m]})
                years.append({'year': year, 'months': month_list})
            res.append({'client': client, 'years': years})
        return res

    def _get_column_mapping(self):
        return [
            ('col_invoice_date', 'Inv. Date', lambda r: r.invoice_date.strftime('%Y-%m-%d') if r.invoice_date else ''),
            ('col_ref_no', 'ref. no (tally)', lambda r: r.ref_no or ''),
            ('col_project_name', 'Project Name', lambda r: r.project_name or ''),
            ('col_party_name', 'Client Name', lambda r: r.partner_id.name or ''),
            ('col_l2_level', 'L2 Level (Client Group)', lambda r: r.l2_level_id.name or ''),
            ('col_l3_level', 'L3 Level (Project)', lambda r: r.l3_level_id.name or ''),
            ('col_invoice_amount', 'Inv. Amount', lambda r: r.invoice_amount or 0.0),
            ('col_ubs_no', 'UBS No.', lambda r: r.ubs_no or ''),
            ('col_contract_lpo_no', 'Contract/LPO No.', lambda r: r.contract_lpo_no or ''),
            ('col_due_day', 'Due Days', lambda r: r.due_days if r.invoice_date else ''),
            ('col_invoice_type', 'Inv. Type', lambda r: r.invoice_type or ''),
            ('col_follow_up', 'Follow up by', lambda r: (
                r.follow_up_user_id.name if r.follow_up_user_id else (r.follow_up or '')
            )),
            ('col_remarks', 'Remarks', lambda r: r.remarks or ''),
            ('col_status', 'Status', lambda r: dict(r._fields['status'].selection).get(r.status, r.status or '')),
        ]

    def _get_active_columns(self):
        column_fields = [
            ('col_invoice_date', 'invoice_date', 'Inv. Date'),
            ('col_ref_no', 'ref_no', 'ref. no (tally)'),
            ('col_project_name', 'project_name', 'Project Name'),
            ('col_party_name', 'partner_id', 'Client Name'),
            ('col_l2_level', 'l2_level_id', 'L2 Level (Client Group)'),
            ('col_l3_level', 'l3_level_id', 'L3 Level (Project)'),
            ('col_invoice_amount', 'invoice_amount', 'Inv. Amount'),
            ('col_ubs_no', 'ubs_no', 'UBS No.'),
            ('col_contract_lpo_no', 'contract_lpo_no', 'Contract/LPO No.'),
            ('col_due_day', 'due_days', 'Due Days'),
            ('col_invoice_type', 'invoice_type', 'Inv. Type'),
            ('col_follow_up', 'follow_up_user_id', 'Follow up by'),
            ('col_remarks', 'remarks', 'Remarks'),
            ('col_status', 'status', 'Status'),
        ]
        res = []
        for col_field, field_name, label in column_fields:
            if self[col_field]:
                res.append({'field_name': field_name, 'label': label})
        return res

    def get_field_val(self, rec, field_name):
        if field_name == 'status':
            return dict(rec._fields['status'].selection).get(rec.status, rec.status or '')
        if field_name == 'due_days':
            return rec.due_days if rec.invoice_date else ''
        if field_name == 'follow_up_user_id':
            if rec.follow_up_user_id:
                return rec.follow_up_user_id.name
            return rec.follow_up or ''
        val = rec[field_name]
        if not val and val != 0:
            return ''
        if hasattr(val, 'name'):
            return val.display_name or val.name or ''
        if hasattr(val, 'strftime'):
            return val.strftime('%Y-%m-%d')
        if isinstance(val, float):
            return val
        return str(val)

    def action_export_excel(self):
        if openpyxl is None:
            raise UserError(_(
                'Python package openpyxl is not installed. '
                'Ask an administrator to install it to export Excel.'
            ))
        records = self._get_filtered_records()
        grouped_data = self._get_grouped_data(records)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'AMC Invoices'

        active_cols = [(h, g) for f, h, g in self._get_column_mapping() if self[f]]
        headers = [c[0] for c in active_cols]

        title_font = Font(name='Calibri', size=16, bold=True)
        sub_font = Font(name='Calibri', size=11, italic=True)
        ws.cell(row=1, column=1, value='AMC INVOICES - BY CLIENT / YEAR / MONTH').font = title_font
        as_of_str = self.as_of_date.strftime('%d %B %Y')
        ws.cell(row=2, column=1, value='as of %s' % as_of_str).font = sub_font

        header_fill = PatternFill(start_color='1A4D2E', end_color='1A4D2E', fill_type='solid')
        header_font_color = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        ws.append([])
        ws.append(['S.No'] + headers)
        for col_idx in range(1, len(headers) + 2):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font_color
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC'),
        )
        client_fill = PatternFill(start_color='C5D5EA', end_color='C5D5EA', fill_type='solid')
        client_font = Font(name='Calibri', size=13, bold=True)
        year_fill = PatternFill(start_color='EBDCB9', end_color='EBDCB9', fill_type='solid')
        year_font = Font(name='Calibri', size=14, bold=True)
        month_fill = PatternFill(start_color='DFBBB4', end_color='DFBBB4', fill_type='solid')
        month_font = Font(name='Calibri', size=11.5, bold=True)

        current_row = 5
        serial_no = 1
        for client_data in grouped_data:
            ws.cell(row=current_row, column=1, value=client_data['client']).font = client_font
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers) + 1)
            for col_idx in range(1, len(headers) + 2):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.fill = client_fill
                cell.alignment = Alignment(vertical='center')
            current_row += 1

            for year_data in client_data['years']:
                ws.cell(row=current_row, column=1, value='YEAR %s' % year_data['year']).font = year_font
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers) + 1)
                for col_idx in range(1, len(headers) + 2):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.fill = year_fill
                    cell.alignment = Alignment(vertical='center')
                current_row += 1

                for month_data in year_data['months']:
                    ws.cell(row=current_row, column=1, value=month_data['name']).font = month_font
                    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers) + 1)
                    for col_idx in range(1, len(headers) + 2):
                        cell = ws.cell(row=current_row, column=col_idx)
                        cell.fill = month_fill
                        cell.alignment = Alignment(vertical='center')
                    current_row += 1

                    for rec in month_data['records']:
                        ws.cell(row=current_row, column=1, value=serial_no).alignment = Alignment(horizontal='center')
                        ws.cell(row=current_row, column=1).border = thin_border
                        serial_no += 1
                        col_idx = 2
                        for header, getter in active_cols:
                            val = getter(rec)
                            cell = ws.cell(row=current_row, column=col_idx, value=val)
                            cell.border = thin_border
                            if header == 'Inv. Amount':
                                cell.number_format = '#,##0.00'
                                cell.alignment = Alignment(horizontal='right')
                            elif header == 'Inv. Date':
                                cell.alignment = Alignment(horizontal='center')
                            col_idx += 1
                        current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        attachment = self.env['ir.attachment'].create({
            'name': 'AMC_Invoices_Statement.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }

    def action_print_pdf(self):
        return self.env.ref('cpabooks_cafm.action_report_amc_invoice_statement').report_action(self)
