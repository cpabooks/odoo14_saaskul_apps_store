# -*- coding: utf-8 -*-

import base64
from datetime import date, datetime, time
import io
import re

from odoo import fields, models, _
from odoo.exceptions import UserError


class CafmAmcRegisterImport(models.TransientModel):
    _name = 'cpabooks.cafm.amc.import'
    _description = 'AMC Register Excel Import'

    file_data = fields.Binary(string='Excel File')
    file_name = fields.Char(string='File Name')

    def _normalize_header(self, value):
        if value is None:
            return ''
        if not isinstance(value, str):
            value = str(value)
        value = value.replace('\ufeff', '').replace('\xa0', ' ').strip().lower()
        value = value.replace('.', ' ').replace('_', ' ').replace('/', ' ')
        value = value.replace('\n', ' ').replace('\r', ' ')
        return re.sub(r'\s+', ' ', value).strip()

    def _column_index(self, header_map, *names):
        for name in names:
            index = header_map.get(self._normalize_header(name))
            if index is not None:
                return index
        return None

    def _cell(self, row, header_map, *names):
        index = self._column_index(header_map, *names)
        if index is None or index >= len(row):
            return False
        return row[index]

    def _to_str(self, value):
        if value in (False, None):
            return False
        if isinstance(value, float) and value == int(value):
            return str(int(value))
        text = str(value).strip()
        return text or False

    def _parse_inv_type(self, value):
        value = (value or '').strip().lower()
        mapping = {
            'y': 'annually',
            'yearly': 'annually',
            'annual': 'annually',
            'annually': 'annually',
            'h': 'half_yearly',
            'half yearly': 'half_yearly',
            'half-yearly': 'half_yearly',
            'halfyearly': 'half_yearly',
            'q': 'quarterly',
            'quarterly': 'quarterly',
            'm': 'monthly',
            'monthly': 'monthly',
        }
        return mapping.get(value, 'quarterly')

    def _parse_invoice_month(self, value):
        """Keep multi tags (Mar, Jun, Sep, Dec); normalize single tokens to Mon abbr."""
        raw = (value or '').strip()
        if not raw:
            return False
        # already multi-tag
        if any(sep in raw for sep in (',', ';', '/', '|')):
            parts = []
            for part in re.split(r'[,;/|]+', raw):
                token = self._parse_invoice_month_token(part)
                if token:
                    parts.append(token)
            return ', '.join(parts) if parts else False
        return self._parse_invoice_month_token(raw)

    def _parse_invoice_month_token(self, value):
        value = (value or '').strip().lower().replace('.', '')
        mapping = {
            '1 jan': 'Jan', 'jan': 'Jan', 'january': 'Jan',
            'feb': 'Feb', 'february': 'Feb',
            'mar': 'Mar', 'march': 'Mar',
            'apr': 'Apr', 'april': 'Apr',
            'may': 'May',
            'jun': 'Jun', 'june': 'Jun',
            'jul': 'Jul', 'july': 'Jul',
            'aug': 'Aug', 'august': 'Aug',
            'sep': 'Sep', 'sept': 'Sep', 'september': 'Sep',
            'oct': 'Oct', 'october': 'Oct',
            'nov': 'Nov', 'november': 'Nov',
            'dec': 'Dec', 'december': 'Dec',
        }
        if not value:
            return False
        return mapping.get(value, (value or '').strip().title() or False)

    def _to_date(self, value):
        """Normalize Excel / openpyxl cell values to datetime.date or False."""
        if value is None or value is False:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, time):
            return False
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if isinstance(value, float) and value != value:  # NaN
                return False
            try:
                from openpyxl.utils.datetime import WINDOWS_EPOCH, from_excel
                dt = from_excel(float(value), WINDOWS_EPOCH)
            except Exception:
                return False
            if isinstance(dt, datetime):
                return dt.date()
            return False
        if isinstance(value, str):
            raw = value.strip()
            if not raw:
                return False
            try:
                return self._to_date(float(raw.replace(',', '')))
            except ValueError:
                pass
            for fmt in (
                '%d-%b-%y', '%d-%b-%Y', '%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d',
                '%d-%m-%y', '%d-%m-%Y', '%d-%b-%y',
            ):
                try:
                    return datetime.strptime(raw, fmt).date()
                except ValueError:
                    continue
        return False

    def _as_date_or_false(self, value):
        if value is None or value is False:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return self._to_date(value)

    def _contract_period_years_from_dates(self, start, end):
        start_d = self._as_date_or_false(start)
        end_d = self._as_date_or_false(end)
        if not start_d or not end_d or end_d < start_d:
            return None
        delta = end_d - start_d
        days = getattr(delta, 'days', None)
        if days is None:
            return None
        return max(1, int(round(days / 365.0)))

    def _date_for_orm(self, value):
        d = self._as_date_or_false(value)
        return fields.Date.to_string(d) if d else False

    def _to_float(self, value):
        if value in (False, None, ''):
            return 0.0
        if isinstance(value, str):
            value = value.replace(',', '').strip()
        try:
            return float(value or 0.0)
        except ValueError:
            return 0.0

    def _to_int(self, value):
        if value in (False, None, ''):
            return False
        try:
            return int(round(self._to_float(value)))
        except (TypeError, ValueError):
            return False

    def _parse_status(self, value):
        value = (value or '').strip().lower()
        if value in ('live', 'active', 'running'):
            return 'active'
        if value in ('pending', 'renewal', 'renewal due'):
            return 'pending'
        if value in ('expired', 'expire'):
            return 'expired'
        if value in ('cancelled', 'canceled', 'cancel'):
            return 'cancelled'
        return 'active'

    def _parse_live_status(self, value):
        value = (value or '').strip().lower()
        if value in ('non live', 'non-live', 'nonlive', 'inactive', 'not live', 'expired', 'cancelled', 'canceled'):
            return 'non_live'
        return 'live'

    def _parse_contract_type(self, value):
        value = (value or '').strip().lower()
        if value == 'var':
            return 'var'
        return 'amc'

    def _parse_invoice_update_status(self, value):
        text = (value or '').strip().lower()
        text = re.sub(r'\s+', ' ', text)
        if not text:
            return False
        if text in ('updated', 'update', 'yes', 'done'):
            return 'updated'
        if text in ('not updated', 'not update', 'pending', 'no'):
            return 'not_updated'
        return False

    def _get_or_create_name(self, model_name, name):
        name = (name or '').strip()
        if not name:
            return False
        record = self.env[model_name].search([('name', '=', name)], limit=1)
        if record:
            return record
        return self.env[model_name].create({'name': name})

    def _get_or_create_project(self, name):
        name = (name or '').strip()
        if not name:
            name = 'AMC Imported Project'
        project = self.env['project.project'].search([('name', '=', name)], limit=1)
        if project:
            return project
        return self.env['project.project'].create({'name': name})

    def _resolve_project_and_unit(self, label):
        label = (label or '').strip()
        if not label:
            return False, False
        unit = self.env['cpabooks.cafm.unit'].search(['|', ('code', '=', label), ('name', '=', label)], limit=1)
        if unit:
            return unit.project_id, unit
        parts = [part.strip() for part in label.split(' - ') if part.strip()]
        project = False
        if parts:
            project = self.env['project.project'].search(['|', ('cafm_code', '=', parts[0]), ('name', '=', parts[0])], limit=1)
        if not project and len(parts) > 1:
            project = self.env['project.project'].search(['|', ('cafm_code', '=', parts[1]), ('name', '=', parts[1])], limit=1)
        if not project:
            project = self.env['project.project'].search([('name', '=', label)], limit=1)
        unit = False
        if project and len(parts) > 2:
            unit_label = parts[-1]
            unit = self.env['cpabooks.cafm.unit'].search([
                ('project_id', '=', project.id),
                '|', ('code', '=', unit_label), ('name', '=', unit_label),
            ], limit=1)
        return project, unit

    # Aliases used for required-column checks and row parsing (keep in sync).
    # Match AMC Tracking sheet headers from register Excel (incl. typos).
    _HDR_PROJECT = (
        'description', 'project name', 'prjoect name', 'proj name', 'project',
        'building', 'property', 'project description', 'tower', 'block', 'community',
    )
    _HDR_PROJECT_LABEL = (
        'project name', 'prjoect name', 'proj name', 'project',
        'building', 'property', 'project description', 'tower', 'block', 'community',
    )
    _HDR_CUSTOMER = (
        'client name', 'customer name', 'customer', 'client',
        'cont', 'contact', 'contact name', 'bill to', 'owner', 'company',
        'company name', 'lessee', 'tenant', 'party', 'account name',
        'cust name', 'cust', 'end user',
    )
    _HDR_INV_FREQ = (
        'inv freq', 'inv frequency', 'invoice frequency', 'invoice freq',
        'inv type', 'freq', 'frequency', 'invoice type', 'billing frequency',
        'billing cycle', 'payment frequency', 'amc frequency',
        'invoice period', 'payment terms', 'billing',
    )
    _HDR_TOTAL_CONTRACT = (
        'total cont value', 'total contract value', 'totl contract value',
        'tot contract value', 'total contract', 'contract value', 'total value',
        'contract amt', 'contract amount', 'tcv',
    )
    _HDR_INVOICE_AMOUNT = (
        'invoice amount', 'invoicing value', 'invoice value',
        'per invoice', 'invoice amt', 'billing amount',
    )
    _HDR_YEARLY = (
        'yearly value', 'yearly', 'annual value', 'yearly contract',
        'yearly contract amount', 'year value',
    )
    _HDR_MONTHLY = (
        'monthly revenue', 'monthly value', 'monthly amount', 'mthly revenue',
    )
    _HDR_BUSINESS_UNIT = (
        'business unit', 'bu', 'bus unit', 'cost centre', 'cost center',
    )
    _HDR_MAJOR_GROUPS = (
        'client group name', 'client group', 'customer group name', 'customer group',
        'major groups', 'major group', 'groups', 'group',
    )
    _HDR_CONT_YEARS = (
        'cont in yr', 'cont in year', 'contract in yr', 'contract years',
        'contract period years', 'period years', 'year', 'years',
    )
    _HDR_LOCATION = (
        'project location', 'location', 'loc', 'area', 'site location',
    )

    def _header_map_from_row(self, row_values):
        headers = [self._normalize_header(value) for value in row_values]
        header_map = {}
        for index, header in enumerate(headers):
            if header and header not in header_map:
                header_map[header] = index
        return header_map

    def _is_business_unit_format(self, header_map):
        return (
            self._column_index(header_map, *self._HDR_BUSINESS_UNIT) is not None
            and self._column_index(header_map, *self._HDR_PROJECT) is not None
        )

    def _missing_required_columns(self, header_map):
        missing = []
        if self._column_index(header_map, *self._HDR_PROJECT) is None:
            missing.append('Project Name / Description / …')
            return missing

        has_value = (
            self._column_index(header_map, *self._HDR_TOTAL_CONTRACT) is not None
            or self._column_index(header_map, *self._HDR_INVOICE_AMOUNT) is not None
            or self._column_index(header_map, *self._HDR_MONTHLY) is not None
            or self._column_index(header_map, *self._HDR_YEARLY) is not None
        )
        if not has_value:
            missing.append('Total Cont. value / Yearly Value / Invoice Amount / Monthly Revenue / …')

        # Classic tracking sheet needs client (+ optional inv freq).
        # Business-unit sheet can omit client (partner created from project / group).
        if not self._is_business_unit_format(header_map):
            if self._column_index(header_map, *self._HDR_CUSTOMER) is None:
                missing.append('Client Name / Customer / …')
        return missing

    def _find_header_row(self, rows, max_scan=15):
        """Use first row that satisfies required columns (handles title rows above headers)."""
        for row_idx in range(min(max_scan, len(rows))):
            header_map = self._header_map_from_row(rows[row_idx])
            if not header_map:
                continue
            if not self._missing_required_columns(header_map):
                return row_idx, header_map
        return None, None

    def action_download_template(self):
        """Excel template matching AMC Register spreadsheet + business-unit sheet."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise UserError(_('openpyxl is required to download the Excel template.'))

        output = io.BytesIO()
        wb = openpyxl.Workbook()

        # Sheet 1 — exact AMC Tracking columns from register Excel
        ws1 = wb.active
        ws1.title = 'AMC Tracking'
        headers1 = [
            'SL', 'Type', 'Prjoect Name', 'Client UBS Number', 'Client Name',
            'Client group name', 'Inv. Freq.', 'Last Invoiced', 'Next Invoice',
            'Inv. Status', 'Inv. Verifying date', 'Project Location', 'Cont. in Yr',
            'Total Cont. value', 'Yearly Value', 'Monthly Revenue', 'Invoice Amount',
            'Cont. Start Date', 'Cont. Expiry Date', 'Cont. Renewal Date',
            'Cont. Status', 'Contract status',
        ]
        ws1.append(headers1)
        header_fill = PatternFill('solid', fgColor='000000')
        header_font = Font(color='FFFF00', bold=True)
        for cell in ws1[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, horizontal='center')
        ws1.append([
            1, 'AMC', 'NHL - 55 Badariya Khalifa Al Mazbali', 'UBS-300', 'ASTECO',
            'ASTECO', 'Q', '31-May-26', 'Aug',
            'Updated', '28-05-26', 'AUH', 1.00,
            11000, 11000, 917, 2750,
            '01-Aug-25', '31/07/2026', '30-Jul-26',
            'Live', 'Auto Renewed - No Cont. Copy',
        ])
        ws1.append([
            2, 'AMC', 'Sh. Ahmed Khalifa Salman Al-Khalifa', '', 'ADCP',
            'ADCP', 'Q', '30-Apr-26', 'July',
            'Not updated', '', 'AUH', 3.00,
            91500, 30500, 2542, 7625,
            '20-Jul-25', '19-Aug-28', '18-Aug-28',
            'Live', 'check & verified',
        ])
        ws1.append([
            3, 'AMC', 'Sample Villa AMC', 'UBS-301', 'PVT',
            'PVT', 'Q', '31-May-26', 'Aug',
            'Updated', '28-05-26', 'AUH', 5.00,
            18225, 3645, 304, 911,
            '01-Aug-25', '31-Aug-30', '30-Aug-30',
            'Live', 'check & verified',
        ])

        # Sheet 2 — business unit / major groups format (alternate upload)
        ws2 = wb.create_sheet('AMC Business Unit')
        headers2 = [
            'SL', 'BUSINESS UNIT', 'Project Name', 'MAJOR GROUPS',
            'Contract Value', 'Monthly Revenue',
        ]
        ws2.append(headers2)
        for cell in ws2[1]:
            cell.fill = header_fill
            cell.font = header_font
        ws2.append([
            1, '1010101', 'Mangrove Village-312 Villa-AMC', 'MANGROVE',
            1799640, 149970,
        ])
        ws2.append([
            2, '1010102', 'Aber Al Mazroui 3 Villas - AMC', 'ABU DHABI',
            64000, 2666.66,
        ])

        wb.save(output)
        output.seek(0)

        attachment = self.env['ir.attachment'].create({
            'name': 'amc_register_import_template.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    def _find_existing_contract(self, amc_model, values, project, unit, partner, description):
        """Prefer UBS / business-unit / project+client match before notes ilike."""
        domain_base = [('project_id', '=', project.id if project else False)]
        if values.get('ubs_no'):
            found = amc_model.search(domain_base + [('ubs_no', '=', values['ubs_no'])], limit=1)
            if found:
                return found
        if values.get('business_unit') and partner:
            found = amc_model.search(domain_base + [
                ('business_unit', '=', values['business_unit']),
                ('client_id', '=', partner.id),
            ], limit=1)
            if found:
                return found
        if partner:
            found = amc_model.search([
                ('project_id', '=', project.id if project else False),
                ('unit_id', '=', unit.id if unit else False),
                ('client_id', '=', partner.id),
            ], limit=1)
            if found:
                return found
        if description:
            found = amc_model.search([
                ('project_id', '=', project.id if project else False),
                ('client_id', '=', partner.id if partner else False),
                ('notes', 'ilike', description),
            ], limit=1)
            if found:
                return found
        return amc_model.browse()

    def action_import_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_('Please select an Excel file to import.'))
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_('openpyxl is required to import Excel AMC files.'))

        workbook = load_workbook(filename=io.BytesIO(base64.b64decode(self.file_data)), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise UserError(_('The Excel file is empty.'))

        header_row_idx, header_map = self._find_header_row(rows)
        if header_map is None:
            lines = []
            for i in range(min(15, len(rows))):
                hm = self._header_map_from_row(rows[i])
                titles = ', '.join(sorted(hm.keys())[:35]) if hm else '(empty or non-text)'
                miss = self._missing_required_columns(hm) if hm else ['(no headers)']
                lines.append('Row %s: %s — still missing: %s' % (i + 1, titles, '; '.join(miss)))
            raise UserError(
                _('No header row found with all required columns in the first %s rows.\n%s') % (
                    15,
                    '\n'.join(lines),
                )
            )

        amc_model = self.env['cpabooks.cafm.contract']
        partner_model = self.env['res.partner']
        created_count = 0
        updated_count = 0
        for row in rows[header_row_idx + 1:]:
            if not any(row):
                continue

            description = self._to_str(self._cell(row, header_map, *self._HDR_PROJECT)) or ''
            contract_type = self._cell(row, header_map, 'type', 'contract type')
            customer_name = self._to_str(self._cell(row, header_map, *self._HDR_CUSTOMER)) or ''
            inv_type = self._cell(row, header_map, *self._HDR_INV_FREQ)
            total_contract = self._cell(row, header_map, *self._HDR_TOTAL_CONTRACT)
            invoice_amount = self._cell(row, header_map, *self._HDR_INVOICE_AMOUNT)
            yearly_value = self._cell(row, header_map, *self._HDR_YEARLY)
            monthly_revenue = self._cell(row, header_map, *self._HDR_MONTHLY)
            start_date = self._to_date(self._cell(
                row,
                header_map,
                'cont start date',
                'start date',
                'contract start date',
                'contract date',
                'starting date',
                'start',
            ))
            expiry_date = self._to_date(self._cell(
                row,
                header_map,
                'cont expiry date',
                'expiry date',
                'contract expiry date',
                'contract expiry',
                'contract end date',
                'end date',
                'ending date',
                'expiry',
                'expire date',
            ))
            renewal_date = self._to_date(self._cell(
                row,
                header_map,
                'cont renewal date',
                'renewal date',
                'contract renewal date',
                'renewal',
            ))
            contract_no = self._cell(row, header_map, 'contract no', 'contract number')
            trn = self._cell(row, header_map, 'trn')
            month = self._cell(
                row, header_map,
                'next invoice', 'next inv', 'next invoice month', 'month',
            )
            group_name = self._cell(row, header_map, *self._HDR_MAJOR_GROUPS)
            year_value = self._cell(row, header_map, *self._HDR_CONT_YEARS)
            # Cont. Status = Live / Non Live (not workflow Status)
            status_value = self._cell(
                row, header_map,
                'cont status', 'live status', 'contract live status',
            )
            # Contract status = check & verified / Auto Renewed - No Cont. Copy
            contract_check = self._to_str(self._cell(
                row, header_map,
                'contract status', 'contract check status', 'check status',
            ))
            location_name = self._to_str(self._cell(row, header_map, *self._HDR_LOCATION))
            units = self._cell(row, header_map, 'units', 'number of flats', 'no of flat / villa')
            update_note = self._cell(row, header_map, 'update', 'cont update')
            remarks = self._cell(row, header_map, 'remarks', 'remark', 'comments', 'comment')
            serial_no = self._to_int(self._cell(row, header_map, 'sl', 's l', 'serial', 'serial no', 'sr'))
            ubs_no = self._to_str(self._cell(
                row, header_map,
                'client ubs number', 'ubs number', 'ubs no', 'ubs', 'client ubs',
            ))
            business_unit = self._to_str(self._cell(row, header_map, *self._HDR_BUSINESS_UNIT))
            last_invoiced = self._to_date(self._cell(
                row, header_map,
                'last invoiced', 'last invoice', 'last inv', 'last invoiced date',
            ))
            invoice_status = self._parse_invoice_update_status(self._cell(
                row, header_map,
                'inv status', 'invoice status', 'invoicing status', 'inv  status',
            ))
            verifying_date = self._to_date(self._cell(
                row, header_map,
                'inv verifying date', 'verifying date', 'verify date',
                'verified date', 'verification date', 'invoice verifying date',
            ))

            project_label = self._cell(row, header_map, *self._HDR_PROJECT_LABEL)
            unit = False
            if project_label:
                project, unit = self._resolve_project_and_unit(str(project_label or ''))
                if not project:
                    project = self._get_or_create_project(str(project_label or description))
            else:
                project = self._get_or_create_project(description)

            partner_name = customer_name
            if not partner_name and group_name:
                partner_name = self._to_str(group_name) or ''
            if not partner_name:
                partner_name = description or 'AMC Import Customer'
            partner = partner_model.search([('name', '=', partner_name)], limit=1)
            if not partner:
                partner = partner_model.create({'name': partner_name})
            if trn:
                partner.vat = str(trn).strip()
            customer_group = self._get_or_create_name(
                'cpabooks.cafm.customer.group',
                self._to_str(group_name) or '',
            )
            location_rec = False
            if location_name:
                location_rec = self._get_or_create_name('cpabooks.cafm.location', location_name)
            notes = []
            if description:
                notes.append(description)
            if trn:
                notes.append('TRN: %s' % trn)
            if month:
                notes.append('Invoice Month: %s' % month)
            if update_note:
                notes.append('Update: %s' % update_note)
            if remarks:
                notes.append('Remarks: %s' % remarks)
            if business_unit:
                notes.append('Business Unit: %s' % business_unit)

            freq = self._parse_inv_type(str(inv_type or '')) if inv_type else 'quarterly'
            total_val = self._to_float(total_contract)
            invoice_amt = self._to_float(invoice_amount) if invoice_amount else 0.0
            yearly_val = self._to_float(yearly_value) if yearly_value else 0.0
            monthly_val = self._to_float(monthly_revenue) if monthly_revenue else 0.0

            # Cont. in Yr wins; else derive from start/expiry dates.
            period = 1
            if year_value not in (False, None, ''):
                period = max(1, min(5, int(round(self._to_float(year_value)))))
            else:
                period_from_dates = self._contract_period_years_from_dates(start_date, expiry_date)
                if period_from_dates is not None:
                    period = period_from_dates

            if not total_val and yearly_val:
                total_val = yearly_val * period
            if not total_val and invoice_amt:
                multiplier = {
                    'monthly': 12,
                    'quarterly': 4,
                    'half_yearly': 2,
                    'annually': 1,
                }.get(freq, 1)
                total_val = invoice_amt * multiplier * period
            if not total_val and monthly_val:
                total_val = monthly_val * 12.0 * period
                if not inv_type:
                    freq = 'monthly'

            live_status = self._parse_live_status(str(status_value or 'live'))
            values = {
                'name': str(contract_no).strip() if contract_no else 'New',
                'amc_contract_type': self._parse_contract_type(str(contract_type or '')),
                'client_id': partner.id,
                'customer_group_id': customer_group.id if customer_group else False,
                'project_id': project.id,
                'unit_id': unit.id if unit else False,
                'invoice_frequency': freq,
                'total_contract_value': total_val,
                'contract_period_years': period,
                'contract_date': self._date_for_orm(start_date),
                'contract_expiry': self._date_for_orm(expiry_date),
                'renewal_date': self._date_for_orm(renewal_date or expiry_date),
                'live_status': live_status,
                'contract_status': 'active' if live_status == 'live' else self._parse_status(str(status_value or '')),
                'number_of_flats': int(self._to_float(units)),
                'remarks': str(remarks or update_note or ''),
                'notes': '\n'.join(notes),
                'serial_no': serial_no or False,
                'ubs_no': ubs_no or False,
                'business_unit': business_unit or False,
                'location': location_name or False,
                'project_location_id': location_rec.id if location_rec else False,
                'last_invoiced_date': self._date_for_orm(last_invoiced),
                'next_invoice_month': self._parse_invoice_month(str(month or '')) if month else False,
                'invoice_update_status': invoice_status or False,
                'contract_check_status': self.env['cpabooks.cafm.contract']._parse_contract_check_status(
                    contract_check
                ) if contract_check else False,
                'verifying_date': self._date_for_orm(verifying_date),
            }
            if invoice_amt and not total_contract:
                # Keep per-invoice amount when Total Cont. value was derived.
                values['invoicing_value'] = invoice_amt

            # Avoid overwriting existing SL / UBS with empty when updating
            clean_values = {k: v for k, v in values.items() if v not in (False, None, '') or k in (
                'unit_id', 'customer_group_id', 'project_location_id', 'contract_date', 'contract_expiry',
                'renewal_date', 'last_invoiced_date', 'verifying_date', 'location',
            )}

            existing = self._find_existing_contract(
                amc_model, values, project, unit, partner, description,
            )
            if existing:
                existing.write(clean_values)
                # Live + expired after import → roll dates so orders can be created later.
                existing._auto_renew_live_if_expired()
                updated_count += 1
            else:
                rec = amc_model.create(values)
                rec._auto_renew_live_if_expired()
                created_count += 1

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('AMC Import'),
                'message': _('%s AMC rows created, %s updated.') % (created_count, updated_count),
                'sticky': False,
                'type': 'success',
            }
        }
