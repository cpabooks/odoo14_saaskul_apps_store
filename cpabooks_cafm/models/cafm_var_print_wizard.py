import base64
import io

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:  # pragma: no cover - tenant images may omit Excel extras
    openpyxl = None
    Font = Alignment = PatternFill = Border = Side = None

from odoo import models, fields, api, _
from odoo.exceptions import UserError

class CafmVarPrintWizard(models.TransientModel):
    _name = 'cpabooks.var.print.wizard'
    _description = 'VAR Work Report Print Wizard'

    partner_id = fields.Many2one('res.partner', string='Client Name')
    as_of_date = fields.Date(string='As of Date', default=fields.Date.context_today, required=True)

    # Column switches (all default to True)
    col_date = fields.Boolean(string='Date', default=True)
    col_party_name = fields.Boolean(string='Client Name', default=True)
    col_call_no = fields.Boolean(string='Call No.', default=True)
    col_call_no_tally = fields.Boolean(string='Call No. (Tally)', default=True)
    col_opened_by = fields.Boolean(string='Open by', default=True)
    col_l2_level = fields.Boolean(string='L2 Level (Client Group)', default=True)
    col_l3_level = fields.Boolean(string='L3 Level (Project)', default=True)
    col_flat_villa = fields.Boolean(string='Flat / Villa', default=True)
    col_qt_no = fields.Boolean(string='QT.No.', default=True)
    col_qt_no_tally = fields.Boolean(string='QT.No. (Tally)', default=True)
    col_qt_date = fields.Boolean(string='QT.Date.', default=True)
    col_work_type = fields.Boolean(string='Type of Work', default=True)
    col_problem_description = fields.Boolean(string='Problem Description', default=True)
    col_work_completion_date = fields.Boolean(string='Wrk.Compl.Date', default=True)
    col_qt_amount = fields.Boolean(string='Qt. Amount', default=True)
    col_work_approved_amount = fields.Boolean(string='Work App. Amt', default=True)
    col_status = fields.Boolean(string='Status', default=True)

    @api.model
    def default_get(self, fields_list):
        res = super(CafmVarPrintWizard, self).default_get(fields_list)
        last_wizard = self.env['cpabooks.var.print.wizard'].search([('create_uid', '=', self.env.uid)], order='id desc', limit=1)
        if last_wizard:
            for col in [f for f in self._fields.keys() if f.startswith('col_')]:
                if col in fields_list:
                    res[col] = getattr(last_wizard, col)
        return res

    def _get_filtered_records(self):
        domain = [('work_date', '<=', self.as_of_date)]
        if self.partner_id:
            domain.append(('partner_id', '=', self.partner_id.id))
        return self.env['cpabooks.cafm.var.work'].search(domain, order='work_date asc')

    def _get_grouped_data(self, records):
        from collections import defaultdict
        grouped = defaultdict(lambda: defaultdict(list))
        for r in records:
            if not r.work_date:
                continue
            year = r.work_date.year
            month = r.work_date.strftime('%B %Y').upper()
            grouped[year][month].append(r)
        
        sorted_years = sorted(grouped.keys())
        res = []
        for y in sorted_years:
            months = grouped[y]
            month_order = {
                'JANUARY': 1, 'FEBRUARY': 2, 'MARCH': 3, 'APRIL': 4,
                'MAY': 5, 'JUNE': 6, 'JULY': 7, 'AUGUST': 8,
                'SEPTEMBER': 9, 'OCTOBER': 10, 'NOVEMBER': 11, 'DECEMBER': 12
            }
            sorted_months = sorted(months.keys(), key=lambda m: month_order.get(m.split()[0], 0))
            month_list = []
            for m in sorted_months:
                month_list.append({
                    'name': m,
                    'records': months[m]
                })
            res.append({
                'year': y,
                'months': month_list
            })
        return res

    def _get_column_mapping(self):
        return [
            ('col_date', 'Date', lambda r: r.work_date.strftime('%Y-%m-%d') if r.work_date else ''),
            ('col_party_name', 'Client Name', lambda r: r.partner_id.name or ''),
            ('col_call_no', 'Call No.', lambda r: r.call_no or ''),
            ('col_call_no_tally', 'Call No. (Tally)', lambda r: r.call_no_tally or ''),
            ('col_opened_by', 'Open by', lambda r: r.opened_by or ''),
            ('col_l2_level', 'L2 Level (Client Group)', lambda r: r.l2_level_id.name or ''),
            ('col_l3_level', 'L3 Level (Project)', lambda r: r.l3_level_id.name or ''),
            ('col_flat_villa', 'Flat / Villa', lambda r: r.flat_villa or ''),
            ('col_qt_no', 'QT.No.', lambda r: r.qt_no or ''),
            ('col_qt_no_tally', 'QT.No. (Tally)', lambda r: r.qt_no_tally or ''),
            ('col_qt_date', 'QT.Date.', lambda r: r.qt_date.strftime('%Y-%m-%d') if r.qt_date else ''),
            ('col_work_type', 'Type of Work', lambda r: r.work_type_id.name or ''),
            ('col_problem_description', 'Problem Description', lambda r: r.problem_description or ''),
            ('col_work_completion_date', 'Wrk.Compl.Date', lambda r: r.work_completion_date.strftime('%Y-%m-%d') if r.work_completion_date else ''),
            ('col_qt_amount', 'Qt. Amount', lambda r: r.qt_amount or 0.0),
            ('col_work_approved_amount', 'Work App. Amt', lambda r: r.work_approved_amount or 0.0),
            ('col_status', 'Status', lambda r: r.status or ''),
        ]

    def _get_active_columns(self):
        column_fields = [
            ('col_date', 'work_date', 'Date'),
            ('col_party_name', 'partner_id', 'Client Name'),
            ('col_call_no', 'call_no', 'Call No.'),
            ('col_call_no_tally', 'call_no_tally', 'Call No. (Tally)'),
            ('col_opened_by', 'opened_by', 'Open by'),
            ('col_l2_level', 'l2_level_id', 'L2 Level (Client Group)'),
            ('col_l3_level', 'l3_level_id', 'L3 Level (Project)'),
            ('col_flat_villa', 'flat_villa', 'Flat / Villa'),
            ('col_qt_no', 'qt_no', 'QT.No.'),
            ('col_qt_no_tally', 'qt_no_tally', 'QT.No. (Tally)'),
            ('col_qt_date', 'qt_date', 'QT.Date.'),
            ('col_work_type', 'work_type_id', 'Type of Work'),
            ('col_problem_description', 'problem_description', 'Problem Description'),
            ('col_work_completion_date', 'work_completion_date', 'Wrk.Compl.Date'),
            ('col_qt_amount', 'qt_amount', 'Qt. Amount'),
            ('col_work_approved_amount', 'work_approved_amount', 'Work App. Amt'),
            ('col_status', 'status', 'Status'),
        ]
        res = []
        for col_field, field_name, label in column_fields:
            if self[col_field]:
                res.append({
                    'field_name': field_name,
                    'label': label
                })
        return res

    def get_field_val(self, rec, field_name):
        val = rec[field_name]
        if not val:
            return ''
        if hasattr(val, 'name'):
            return val.display_name or val.name or ''
        if isinstance(val, fields.Date) or hasattr(val, 'strftime'):
            return val.strftime('%Y-%m-%d')
        if isinstance(val, float):
            return val
        return str(val)

    def action_export_excel(self):
        if openpyxl is None:
            raise UserError(_(
                'Python package openpyxl is not installed. '
                'Ask an administrator to install it to export Excel.'
            ))
        records = self._get_filtered_records()
        grouped_data = self._get_grouped_data(records)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "VAR Work Statement"

        active_cols = []
        for field, header, getter in self._get_column_mapping():
            if self[field]:
                active_cols.append((header, getter))

        title_font = Font(name='Calibri', size=16, bold=True)
        sub_font = Font(name='Calibri', size=11, italic=True)
        
        ws.cell(row=1, column=1, value="VAR WORK STATEMENT - BY YEAR/MONTH").font = title_font
        as_of_str = self.as_of_date.strftime('%d %B %Y')
        ws.cell(row=2, column=1, value=f"as of {as_of_str}").font = sub_font
        
        header_fill = PatternFill(start_color="1A4D2E", end_color="1A4D2E", fill_type="solid")
        header_font_color = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
        
        headers = [col[0] for col in active_cols]
        ws.append([]) # row 3 empty
        
        ws.append(['S.No'] + headers)
        ws.row_dimensions[4].height = 25
        for col_idx in range(1, len(headers) + 2):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font_color
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )

        year_fill = PatternFill(start_color="EBDCB9", end_color="EBDCB9", fill_type="solid")
        year_font = Font(name='Calibri', size=14, bold=True)
        
        month_fill = PatternFill(start_color="DFBBB4", end_color="DFBBB4", fill_type="solid")
        month_font = Font(name='Calibri', size=11.5, bold=True)

        current_row = 5
        serial_no = 1

        for year_data in grouped_data:
            ws.cell(row=current_row, column=1, value=f"YEAR {year_data['year']}").font = year_font
            ws.row_dimensions[current_row].height = 28
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers) + 1)
            for col_idx in range(1, len(headers) + 2):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.fill = year_fill
                cell.alignment = Alignment(vertical="center")
            current_row += 1

            for month_data in year_data['months']:
                ws.cell(row=current_row, column=1, value=month_data['name']).font = month_font
                ws.row_dimensions[current_row].height = 22
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers) + 1)
                for col_idx in range(1, len(headers) + 2):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.fill = month_fill
                    cell.alignment = Alignment(vertical="center")
                current_row += 1

                for rec in month_data['records']:
                    ws.cell(row=current_row, column=1, value=serial_no).alignment = Alignment(horizontal="center")
                    ws.cell(row=current_row, column=1).border = thin_border
                    serial_no += 1
                    
                    col_idx = 2
                    for header, getter in active_cols:
                        val = getter(rec)
                        cell = ws.cell(row=current_row, column=col_idx, value=val)
                        cell.border = thin_border
                        
                        if header in ('Qt. Amount', 'Work Appr. Amt', 'Work App. Amt'):
                            cell.number_format = '#,##0.00'
                            cell.alignment = Alignment(horizontal="right")
                        elif header in ('Date', 'QT.Date.', 'Wrk.Compl.Date'):
                            cell.alignment = Alignment(horizontal="center")
                        col_idx += 1
                    current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        attachment = self.env['ir.attachment'].create({
            'name': 'VAR_Work_Statement.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': 'cpabooks.var.print.wizard',
            'res_id': self.id,
        })
        
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_print_pdf(self):
        return self.env.ref('cpabooks_cafm.action_report_var_followup').report_action(self)
