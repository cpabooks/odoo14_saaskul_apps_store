# -*- coding: utf-8 -*-

import base64
from datetime import date, datetime, time
import io
import re

from odoo import fields, models, _
from odoo.exceptions import UserError


class CafmAmcInvoiceImport(models.TransientModel):
    _name = 'cpabooks.cafm.amc.invoice.import'
    _description = 'AMC Invoice Excel Import'

    file_data = fields.Binary(string='Excel File')
    file_name = fields.Char(string='File Name')

    def _normalize_header(self, value):
        if value is None:
            return ''
        if not isinstance(value, str):
            value = str(value)
        value = value.replace('\ufeff', '').replace('\xa0', ' ').strip().lower()
        value = value.replace('.', ' ').replace('_', ' ').replace('/', ' ')
        value = value.replace('\n', ' ').replace('\r', ' ')
        return re.sub(r'\s+', ' ', value).strip()

    def _column_index(self, header_map, *names):
        for name in names:
            index = header_map.get(self._normalize_header(name))
            if index is not None:
                return index
        return None

    def _cell(self, row, header_map, *names):
        index = self._column_index(header_map, *names)
        if index is None or index >= len(row):
            return False
        return row[index]

    def _to_date(self, value):
        if value is None or value is False:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, time):
            return False
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if isinstance(value, float) and value != value:
                return False
            try:
                from openpyxl.utils.datetime import WINDOWS_EPOCH, from_excel
                dt = from_excel(float(value), WINDOWS_EPOCH)
            except Exception:
                return False
            if isinstance(dt, datetime):
                return dt.date()
            return False
        if isinstance(value, str):
            raw = value.strip()
            if not raw:
                return False
            try:
                return self._to_date(float(raw.replace(',', '')))
            except ValueError:
                pass
            for fmt in (
                '%d-%b-%y', '%d-%b-%Y', '%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d',
                '%d-%m-%y', '%d-%m-%Y', '%d/%m/%Y',
            ):
                try:
                    return datetime.strptime(raw, fmt).date()
                except ValueError:
                    continue
        return False

    def _date_for_orm(self, value):
        d = self._to_date(value)
        return fields.Date.to_string(d) if d else False

    def _to_float(self, value):
        if value in (False, None, ''):
            return 0.0
        if isinstance(value, str):
            value = value.replace(',', '').strip()
        try:
            return float(value or 0.0)
        except ValueError:
            return 0.0

    def _to_str(self, value):
        if value in (False, None):
            return False
        if isinstance(value, float) and value == int(value):
            return str(int(value))
        text = str(value).strip()
        return text or False

    def _normalize_status_text(self, value):
        if value in (False, None):
            return ''
        text = str(value).strip().lower()
        text = re.sub(r'\s+', ' ', text)
        return text

    def _map_status_token(self, text):
        text = self._normalize_status_text(text)
        if not text:
            return False
        mapping = {
            'waiting for kpi': 'waiting_kpi',
            'waiting kpi': 'waiting_kpi',
            'kpi': 'waiting_kpi',
            'waiting for csr': 'waiting_csr',
            'waiting csr': 'waiting_csr',
            'csr': 'waiting_csr',
            'not submitted': 'not_submitted',
            'invoice not submitted': 'not_submitted',
            'invoices not submitted': 'not_submitted',
            'not submit': 'not_submitted',
            'submitted': 'submitted',
            'invoice submitted': 'submitted',
        }
        return mapping.get(text, False)

    def _parse_status_from_row(self, row, header_map):
        """Map Excel Status / Invoice Status / KPI Status / CSR Status → one Selection."""
        single = self._map_status_token(
            self._cell(row, header_map, 'status', 'invoice status', 'inv status')
        )
        if single:
            return single

        kpi = self._normalize_status_text(
            self._cell(row, header_map, 'kpi status', 'kpi')
        )
        csr = self._normalize_status_text(
            self._cell(row, header_map, 'csr status', 'csr')
        )
        inv = self._normalize_status_text(
            self._cell(row, header_map, 'invoice status', 'inv status')
        )

        # Prefer active waiting buckets, then not submitted / submitted.
        for blob in (kpi, csr, inv):
            mapped = self._map_status_token(blob)
            if mapped in ('waiting_kpi', 'waiting_csr'):
                return mapped
        for blob in (inv, kpi, csr):
            mapped = self._map_status_token(blob)
            if mapped:
                return mapped
        return 'not_submitted'

    def action_download_template(self):
        output = io.BytesIO()
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'AMC Invoices'

        headers = [
            'Inv. Date', 'ref. no (tally)', 'Project Name', 'Inv. Amount', 'UBS No.',
            'Contract/LPO No.', 'Due Day', 'Inv. Type', 'Follow up', 'Remarks',
            'Status', 'L2 Level (Client Group)', 'L3 Level (Project)', 'Client Name',
        ]
        ws.append(headers)
        ws.append([
            '28-Feb-2026', '26G0089A', '163 Found Property Building', 11628.75, '163',
            '7382', '', 'AMC', 'Nihil', 'Reminder sent on 22/07/2026',
            'Waiting for KPI', '', '163 Found Property Building', '163 Found Property Building',
        ])
        ws.append([
            '30-Apr-2026', '26G0154A', '164 Found Property Holding', 3937.50, '164',
            '7383', '', 'AMC', 'Ragesh', '',
            'Not Submitted', '', '164 Found Property Holding', '164 Found Property Holding',
        ])
        wb.save(output)
        output.seek(0)

        attachment = self.env['ir.attachment'].create({
            'name': 'amc_invoice_import_template.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    def _resolve_l2_l3_partner(self, Invoice, row, header_map, project_name):
        """Same logic as VAR: get/create L2, L3, Client Name."""
        l2_val = self._to_str(self._cell(
            row, header_map,
            'l2 level (client group)', 'l2 level', 'l2 level (project group)', 'client group',
        ))
        l3_val = self._to_str(self._cell(
            row, header_map,
            'l3 level (project)', 'l3 level',
        )) or project_name
        client_val = self._to_str(self._cell(
            row, header_map,
            'client name', 'party name', 'party name (client)', 'customer', 'partner',
        ))
        if not client_val and l3_val:
            client_val = l3_val

        l2 = Invoice._get_or_create_l2_from_name(l2_val) if l2_val else False
        l3 = Invoice._get_or_create_l3_from_name(l3_val) if l3_val else False
        partner = Invoice._get_or_create_partner_from_name(client_val) if client_val else False
        return {
            'l2_level_id': l2.id if l2 else False,
            'l3_level_id': l3.id if l3 else False,
            'partner_id': partner.id if partner else False,
            'project_name': project_name or (l3.name if l3 else False),
        }

    def action_import_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_('Please select an Excel file to import.'))

        import openpyxl
        try:
            wb = openpyxl.load_workbook(
                filename=io.BytesIO(base64.b64decode(self.file_data)),
                data_only=True,
            )
        except Exception as exc:
            raise UserError(_('Could not read Excel file: %s') % exc) from exc

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise UserError(_('The Excel file is empty.'))

        header_row_idx = None
        header_map = {}
        for idx, row in enumerate(rows[:15]):
            cells = [self._normalize_header(c) for c in row]
            if 'ref no' in cells or 'inv date' in cells or 'project name' in cells:
                header_row_idx = idx
                header_map = {
                    self._normalize_header(cell): col
                    for col, cell in enumerate(row)
                    if cell not in (None, False, '')
                }
                break
        if header_row_idx is None:
            raise UserError(_(
                'Could not find header row. Expected columns like '
                'Inv. Date, Ref. No., Project Name, Status.'
            ))

        Invoice = self.env['cpabooks.cafm.amc.invoice.reg']
        created = updated = skipped = 0
        for row in rows[header_row_idx + 1:]:
            if not row or not any(row):
                continue
            ref_no = self._to_str(self._cell(
                row, header_map,
                'ref no', 'ref. no', 'ref no (tally)', 'ref. no (tally)',
                'reference', 'invoice no', 'inv no',
            ))
            project_name = self._to_str(self._cell(
                row, header_map, 'project name', 'project',
            ))
            invoice_date = self._date_for_orm(self._cell(
                row, header_map, 'inv date', 'invoice date', 'inv. date',
            ))
            if not ref_no and not project_name and not invoice_date:
                skipped += 1
                continue

            related = self._resolve_l2_l3_partner(Invoice, row, header_map, project_name)
            vals = {
                'invoice_date': invoice_date,
                'ref_no': ref_no,
                'project_name': related['project_name'],
                'l2_level_id': related['l2_level_id'],
                'l3_level_id': related['l3_level_id'],
                'partner_id': related['partner_id'],
                'invoice_amount': self._to_float(self._cell(
                    row, header_map, 'inv amount', 'invoice amount', 'inv. amount', 'amount',
                )),
                'ubs_no': self._to_str(self._cell(row, header_map, 'ubs no', 'ubs. no', 'ubs')),
                'contract_lpo_no': self._to_str(self._cell(
                    row, header_map,
                    'contract lpo no', 'contract lpo', 'lpo no', 'contract no',
                )),
                'due_day': self._to_str(self._cell(row, header_map, 'due day', 'due date', 'due')),
                'invoice_type': self._to_str(self._cell(
                    row, header_map, 'inv type', 'invoice type', 'inv. type',
                )) or 'AMC',
                'follow_up': self._to_str(self._cell(
                    row, header_map, 'follow up', 'followup', 'followed by',
                )),
                'remarks': self._to_str(self._cell(row, header_map, 'remarks', 'remark', 'note')),
                'status': self._parse_status_from_row(row, header_map),
            }

            existing = Invoice.browse()
            if ref_no:
                existing = Invoice.search([('ref_no', '=', ref_no)], limit=1)
            if existing:
                existing.write(vals)
                updated += 1
            else:
                Invoice.create(vals)
                created += 1

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('AMC Invoice Import'),
                'message': _(
                    'Created: %(created)s, Updated: %(updated)s, Skipped: %(skipped)s'
                ) % {
                    'created': created,
                    'updated': updated,
                    'skipped': skipped,
                },
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }
