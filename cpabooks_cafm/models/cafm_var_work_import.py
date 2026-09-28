# -*- coding: utf-8 -*-

import base64
from datetime import date, datetime, time
import io
import re

from odoo import fields, models, api, _
from odoo.exceptions import UserError


class CafmVarWorkImport(models.TransientModel):
    _name = 'cpabooks.cafm.var.import'
    _description = 'VAR Work Excel Import'

    file_data = fields.Binary(string='Excel File')
    file_name = fields.Char(string='File Name')
    missing_preview = fields.Text(string='Missing related records', readonly=True)
    allow_create_missing = fields.Boolean(
        string='Create missing L2/L3/etc. and continue',
        default=False,
    )

    def _normalize_header(self, value):
        if value is None:
            return ''
        if not isinstance(value, str):
            value = str(value)
        value = value.replace('\ufeff', '').replace('\xa0', ' ').strip().lower()
        value = value.replace('.', ' ').replace('_', ' ').replace('/', ' ')
        value = value.replace('\n', ' ').replace('\r', ' ')
        return re.sub(r'\s+', ' ', value).strip()

    def _column_index(self, header_map, *names):
        for name in names:
            index = header_map.get(self._normalize_header(name))
            if index is not None:
                return index
        return None

    def _cell(self, row, header_map, *names):
        index = self._column_index(header_map, *names)
        if index is None or index >= len(row):
            return False
        return row[index]

    def _to_date(self, value):
        if value is None or value is False:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, time):
            return False
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if isinstance(value, float) and value != value:  # NaN
                return False
            try:
                from openpyxl.utils.datetime import WINDOWS_EPOCH, from_excel
                dt = from_excel(float(value), WINDOWS_EPOCH)
            except Exception:
                return False
            if isinstance(dt, datetime):
                return dt.date()
            return False
        if isinstance(value, str):
            raw = value.strip()
            if not raw:
                return False
            try:
                return self._to_date(float(raw.replace(',', '')))
            except ValueError:
                pass
            for fmt in (
                '%d-%b-%y', '%d-%b-%Y', '%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d',
                '%d-%m-%y', '%d-%m-%Y',
            ):
                try:
                    return datetime.strptime(raw, fmt).date()
                except ValueError:
                    continue
        return False

    def _as_date_or_false(self, value):
        if value is None or value is False:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return self._to_date(value)

    def _date_for_orm(self, value):
        d = self._as_date_or_false(value)
        return fields.Date.to_string(d) if d else False

    def _to_float(self, value):
        if value in (False, None, ''):
            return 0.0
        if isinstance(value, str):
            value = value.replace(',', '').strip()
        try:
            return float(value or 0.0)
        except ValueError:
            return 0.0

    def _parse_status(self, value):
        value = (value or '').strip().lower()
        mapping = {
            'waiting for work approval': 'waiting_approval',
            'waiting approval': 'waiting_approval',
            'var work approval': 'waiting_approval',
            'work ongoing': 'work_ongoing',
            'work on going': 'work_ongoing',
            'ongoing': 'work_ongoing',
            'waiting for lpo': 'waiting_lpo',
            'waiting lpo': 'waiting_lpo',
            'waiting for report completion': 'waiting_report',
            'waiting report': 'waiting_report',
            'var work - invoiced': 'invoiced',
            'var work invoiced': 'invoiced',
            'invoiced': 'invoiced',
            'no status found': 'no_status',
            'status not available': 'no_status',
            'closed': 'closed',
            'cancelled': 'cancelled',
            'canceled': 'cancelled',
        }
        return mapping.get(value, 'no_status')

    def _get_or_create_partner(self, name):
        name = (name or '').strip()
        if not name:
            return False
        partner = self.env['res.partner'].search([('name', '=', name)], limit=1)
        if partner:
            return partner
        return self.env['res.partner'].create({'name': name})

    def _get_or_create_work_type(self, name):
        name = (name or '').strip()
        if not name:
            return False
        work_type = self.env['cpabooks.cafm.work.type'].search([('name', '=', name)], limit=1)
        if work_type:
            return work_type
        return self.env['cpabooks.cafm.work.type'].create({'name': name})

    def _get_or_create_technician(self, name):
        name = (name or '').strip()
        if not name:
            return False
        tech = self.env['cpabooks.cafm.technician'].search([('name', '=', name)], limit=1)
        if tech:
            return tech
        return self.env['cpabooks.cafm.technician'].create({'name': name})

    def _find_only(self, model_name, name):
        name = (name or '').strip()
        if not name:
            return self.env[model_name]
        return self.env[model_name].search([('name', '=', name)], limit=1)

    def _scan_missing_related(self, data_rows, header_map):
        """Collect unseen L2 / L3 / partner / work type names (no create yet)."""
        missing = {
            'L2 Level': set(),
            'L3 Level': set(),
            'Client Name': set(),
            'Type of Work': set(),
            'Technician': set(),
        }
        for row in data_rows:
            if not any(row):
                continue
            l2 = self._cell(row, header_map, 'l2 level', 'l2 level (project group)', 'project group')
            l3 = self._cell(row, header_map, 'l3 level', 'l3 level (project)', 'project')
            party = self._cell(
                row, header_map,
                'party name', 'party name (client)', 'customer', 'partner', 'client name',
            )
            wtype = self._cell(row, header_map, 'type of work', 'work type')
            tech = self._cell(row, header_map, 'technician', 'tech')
            if l2 and not self._find_only('cpabooks.cafm.customer.group', str(l2)):
                missing['L2 Level'].add(str(l2).strip())
            if l3 and not self._find_only('project.project', str(l3)):
                missing['L3 Level'].add(str(l3).strip())
            if party and not self._find_only('res.partner', str(party)):
                missing['Client Name'].add(str(party).strip())
            if wtype and not self._find_only('cpabooks.cafm.work.type', str(wtype)):
                missing['Type of Work'].add(str(wtype).strip())
            if tech and not self._find_only('cpabooks.cafm.technician', str(tech)):
                missing['Technician'].add(str(tech).strip())
        return {k: sorted(v, key=lambda n: n.lower()) for k, v in missing.items() if v}

    def _format_missing_message(self, missing):
        blocks = []
        for label, names in missing.items():
            shown = names[:40]
            block = '%s (%d):\n- %s' % (label, len(names), '\n- '.join(shown))
            if len(names) > 40:
                block += '\n- ... (+%d more)' % (len(names) - 40)
            blocks.append(block)
        return (
            "These records were not found and need to be created to continue import:\n\n"
            "%s\n\n"
            "Tick «Create missing…» below and click Upload & Import again, or Cancel."
        ) % '\n\n'.join(blocks)

    def _get_or_create_l2_level(self, name):
        name = (name or '').strip()
        if not name:
            return False
        group = self.env['cpabooks.cafm.customer.group'].search([('name', '=', name)], limit=1)
        if group:
            return group
        return self.env['cpabooks.cafm.customer.group'].create({'name': name})

    def _get_or_create_l3_level(self, name):
        name = (name or '').strip()
        if not name:
            return False
        project = self.env['project.project'].search([('name', '=', name)], limit=1)
        if project:
            return project
        return self.env['project.project'].create({'name': name})

    def _get_user_by_name(self, name):
        name = (name or '').strip()
        if not name:
            return False
        user = self.env['res.users'].search([('name', '=ilike', name)], limit=1)
        if not user:
            user = self.env['res.users'].search([('login', '=ilike', name)], limit=1)
        return user

    def action_download_template(self):
        output = io.BytesIO()
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "VAR Work Template"

        headers = [
            'SLNo', 'QT.No (Tally)', 'QT.Date', 'Call No. (Tally)', 'Call Dt.', 'Work Appr.By.',
            'Work Appr.Dt.', 'Work App. Amt', 'LPO No.', 'LPO Dt.', 'LPO Amount',
            'Work Comp.Dt.', 'Report Date', 'Inv. No.', 'Inv. Dt.', 'Inv. Amt',
            'Open By', 'L2 Level (Client Group)', 'L3 Level (Project)', 'Flat / Villa', 'Problem Description',
            'Status'
        ]
        ws.append(headers)

        sample_row = [
            1, '24G00001Q', '02-01-2024', '24G00001V', '02-01-2024', 'The Manager',
            '02-01-2024', 550.00, '14759', '27-03-2024', 550.00,
            '06-01-2024', '06-01-2024', '24G0133V', '30-03-2024', 550.00,
            'ragesh', 'MBZ CITY GROUP', 'NHL 13 Al Rumaithi', 'Villa 6', 'Supply and Replace',
            'VAR Work Approval'
        ]
        ws.append(sample_row)

        wb.save(output)
        output.seek(0)

        attachment = self.env['ir.attachment'].create({
            'name': 'var_work_import_template.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': 'cpabooks.cafm.var.import',
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }


    def action_import_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError('Please select an Excel file to upload first.')
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError('openpyxl is required to import Excel files.')

        workbook = load_workbook(filename=io.BytesIO(base64.b64decode(self.file_data)), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise UserError('The Excel file is empty.')

        # Find the header row dynamically
        header_row_index = -1
        header_map = {}
        for index, row in enumerate(rows):
            temp_map = {}
            for col_idx, cell in enumerate(row):
                normalized = self._normalize_header(cell)
                if normalized:
                    temp_map[normalized] = col_idx
            
            # Check if this row looks like the header row (contains tally QT or plain QT)
            if (
                'qt no tally' in temp_map
                or 'qt no (tally)' in temp_map
                or 'qt no' in temp_map
                or 'qt number' in temp_map
                or 'quotation no' in temp_map
            ):
                header_row_index = index
                header_map = temp_map
                break

        if header_row_index == -1:
            # Fallback to row 0 if we couldn't find a row with 'qt no'
            headers_row = rows[0]
            header_map = {}
            for col_idx, header in enumerate(headers_row):
                normalized = self._normalize_header(header)
                if normalized and normalized not in header_map:
                    header_map[normalized] = col_idx
            data_rows = rows[1:]
        else:
            data_rows = rows[header_row_index + 1:]

        # Check required columns — upload needs Tally QT (or legacy QT.No)
        if (
            self._column_index(header_map, 'qt no tally', 'qt no (tally)') is None
            and self._column_index(header_map, 'qt no', 'qt number', 'quotation no') is None
        ):
            raise UserError(
                'Missing required column: QT.No (Tally) (or QT.No for legacy files).'
            )

        missing = self._scan_missing_related(data_rows, header_map)
        if missing and not self.allow_create_missing:
            self.write({
                'missing_preview': self._format_missing_message(missing),
                'allow_create_missing': False,
            })
            return {
                'type': 'ir.actions.act_window',
                'name': 'Import VAR Work — confirm create',
                'res_model': self._name,
                'res_id': self.id,
                'view_mode': 'form',
                'target': 'new',
                'context': dict(self.env.context),
            }

        var_work_model = self.env['cpabooks.cafm.var.work']
        created_count = 0
        updated_count = 0

        for row in data_rows:
            if not any(row):
                continue

            work_date_val = self._cell(row, header_map, 'date', 'work date', 'qt date', 'qt.date', 'call dt.', 'call dt')
            party_name_val = self._cell(
                row, header_map,
                'party name', 'party name (client)', 'customer', 'partner', 'client name',
            )
            call_no_val = self._cell(row, header_map, 'call no', 'call number', 'call no.')
            call_no_tally_val = self._cell(
                row, header_map,
                'call no tally', 'call no (tally)', 'call no.(tally)',
            )
            open_by_val = self._cell(row, header_map, 'open by')
            l2_level_val = self._cell(
                row, header_map,
                'l2 level', 'l2 level (client group)', 'l2 level (project group)', 'project group', 'client group',
            )
            l3_level_val = self._cell(
                row, header_map, 'l3 level', 'l3 level (project)', 'project',
            )
            flat_villa_val = self._cell(row, header_map, 'flat villa', 'villa', 'flat', 'flat / villa')
            qt_no_val = self._cell(row, header_map, 'qt no', 'qt number', 'quotation no', 'qt.no')
            qt_no_tally_val = self._cell(
                row, header_map,
                'qt no tally', 'qt no (tally)', 'qt.no (tally)', 'qt.no tally',
            )
            qt_date_val = self._cell(row, header_map, 'qt date', 'quotation date', 'qt.date')
            call_dt_val = self._cell(row, header_map, 'call dt', 'call dt.', 'call date')
            work_type_val = self._cell(row, header_map, 'type of work', 'work type')
            problem_val = self._cell(row, header_map, 'problem description', 'description', 'problem')
            completion_date_val = self._cell(
                row, header_map,
                'wrk compl date', 'work comp dt', 'work completion date', 'completion date',
                'work comp.dt.', 'work comp.dt',
            )
            qt_amount_val = self._cell(row, header_map, 'qt amount', 'qt amt', 'quotation amount')
            approved_amount_val = self._cell(
                row, header_map,
                'work app amt', 'work appr amt', 'work appr. amt', 'work app. amt',
                'work approved amount', 'approved amount', 'work app.',
            )
            status_val = self._cell(row, header_map, 'status')

            work_approved_by_val = self._cell(
                row, header_map, 'work appr by', 'work appr.by', 'work approved by', 'work appr.by.',
            )
            work_approval_date_val = self._cell(
                row, header_map, 'work appr dt', 'work appr.dt', 'work approval date', 'work appr.dt.',
            )

            lpo_no_val = self._cell(row, header_map, 'lpo no', 'lpo number', 'lpo no.')
            lpo_date_val = self._cell(row, header_map, 'lpo dt', 'lpo date', 'lpo dt.')
            lpo_amount_val = self._cell(row, header_map, 'lpo amount', 'lpo amt')
            report_date_val = self._cell(row, header_map, 'report date')
            tech_val = self._cell(row, header_map, 'technician', 'tech')
            invoice_no_val = self._cell(row, header_map, 'inv no', 'inv. no', 'invoice no', 'invoice number')
            invoice_date_val = self._cell(row, header_map, 'inv dt', 'inv. dt', 'invoice date')
            invoice_amt_val = self._cell(row, header_map, 'inv amt', 'inv. amt', 'invoice amount')

            call = False
            # Prefer Tally call number from Excel upload, then plain Call No.
            lookup_call = call_no_tally_val or call_no_val
            if lookup_call:
                call = self.env['maintenance.request'].search(
                    [('call_no', '=', str(lookup_call).strip())], limit=1,
                )
                if not call:
                    call = self.env['maintenance.request'].search(
                        [('call_no_tally', '=', str(lookup_call).strip())], limit=1,
                    )

            if not party_name_val and call and call.partner_id:
                party_name_val = call.partner_id.name
            if not party_name_val and l3_level_val:
                party_name_val = l3_level_val
            if not party_name_val and call and call.contract_id and call.contract_id.client_id:
                party_name_val = call.contract_id.client_id.name

            if not work_date_val:
                if qt_date_val:
                    work_date_val = qt_date_val
                elif call and call.call_date:
                    work_date_val = call.call_date
                else:
                    work_date_val = date.today()

            partner = self._get_or_create_partner(party_name_val)
            work_type = self._get_or_create_work_type(work_type_val)
            tech = self._get_or_create_technician(tech_val)
            l2_level = self._get_or_create_l2_level(l2_level_val)
            l3_level = self._get_or_create_l3_level(l3_level_val)

            values = {
                'work_date': self._date_for_orm(work_date_val),
                'partner_id': partner.id if partner else False,
                'flat_villa': str(flat_villa_val or '').strip(),
                'work_type_id': work_type.id if work_type else False,
                'problem_description': str(problem_val or '').strip(),
                'qt_amount': self._to_float(qt_amount_val),
                'work_approved_amount': self._to_float(approved_amount_val),
                'lpo_no': str(lpo_no_val or '').strip(),
                'lpo_date': self._date_for_orm(lpo_date_val),
                'lpo_amount': self._to_float(lpo_amount_val),
                'work_completion_date': self._date_for_orm(completion_date_val),
                'report_date': self._date_for_orm(report_date_val),
                'technician_id': tech.id if tech else False,
                'status': self._parse_status(status_val),
                'opened_by': str(open_by_val or '').strip(),
                'qt_date': self._date_for_orm(qt_date_val),
                'call_dt': self._date_for_orm(call_dt_val),
                'l2_level_id': l2_level.id if l2_level else False,
                'l3_level_id': l3_level.id if l3_level else False,
                'work_approved_by': str(work_approved_by_val or '').strip(),
                'work_approval_date': self._date_for_orm(work_approval_date_val),
                'invoice_no': str(invoice_no_val or '').strip() or False,
                'invoice_date': self._date_for_orm(invoice_date_val),
                'invoice_amount': self._to_float(invoice_amt_val),
            }

            if call:
                values['call_id'] = call.id
                if call.call_no:
                    values['call_no'] = call.call_no
                if call_no_tally_val:
                    values['call_no_tally'] = str(call_no_tally_val).strip()
                elif call.call_no_tally:
                    values['call_no_tally'] = call.call_no_tally
            else:
                if call_no_tally_val:
                    values['call_no_tally'] = str(call_no_tally_val).strip()
                if call_no_val:
                    values['call_no'] = str(call_no_val).strip()

            if qt_no_tally_val:
                values['qt_no_tally'] = str(qt_no_tally_val).strip()
            # Legacy Excel without "(Tally)" header → treat QT.No column as Tally upload
            elif qt_no_val and not qt_no_tally_val:
                values['qt_no_tally'] = str(qt_no_val).strip()
            if qt_no_val and qt_no_tally_val:
                # Both columns present: Odoo QT.No + Tally
                values['qt_no'] = str(qt_no_val).strip()

            if invoice_no_val:
                inv = self.env['account.move'].search([
                    ('name', '=', str(invoice_no_val).strip()),
                    ('move_type', 'in', ('out_invoice', 'out_refund'))
                ], limit=1)
                if inv:
                    values['invoice_id'] = inv.id

            # Match existing by Tally QT first, then Odoo QT.No
            existing = False
            if values.get('qt_no_tally'):
                existing = var_work_model.search(
                    [('qt_no_tally', '=', values['qt_no_tally'])], limit=1,
                )
            if not existing and values.get('qt_no'):
                existing = var_work_model.search([('qt_no', '=', values['qt_no'])], limit=1)
            if not existing and values['work_date'] and values['partner_id'] and values['flat_villa']:
                existing = var_work_model.search([
                    ('work_date', '=', values['work_date']),
                    ('partner_id', '=', values['partner_id']),
                    ('flat_villa', '=', values['flat_villa']),
                ], limit=1)

            if existing:
                existing.write(values)
                updated_count += 1
            else:
                var_work_model.create(values)
                created_count += 1

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'VAR Work Import',
                'message': '%s VAR Work rows created, %s updated.' % (created_count, updated_count),
                'sticky': False,
                'type': 'success',
            }
        }
